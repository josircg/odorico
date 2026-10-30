from openpyxl import load_workbook
from validate_email import validate_email
from django.core.exceptions import ValidationError

from .models import Discente, Inscricao, Contato, ListaContato, ListaInteresse, Curso, CursoDocente


def import_pauta_xlsx(curso, tipo, arquivo_xlsx):
    novos_discentes = 0
    tot_lidos = 0
    tot_invalidos = 0
    erros = []
    if not curso:
        raise ValidationError('Curso não definido')

    try:
        wb = load_workbook(arquivo_xlsx)
        ws = wb.active

        for row in ws.iter_rows(min_row=1, values_only=True):
            if len(row) < 1:
                if tot_lidos == 0:
                    erros.append('Arquivo vazio ou com linhas em branco')
                break

            tot_lidos += 1
            nome = row[0]
            if nome:
                nome = nome.strip()
            else:
                erros.append(f"Linha {tot_lidos}: O nome não está preenchido")
                tot_invalidos += 1
                break

            email = row[1]
            if email:
                email = email.strip()
                try:
                    is_valid = validate_email(email, smtp_timeout=10)
                except ValidationError:
                    is_valid = False
            else:
                is_valid = False

            if not is_valid:
                erros.append(f"Linha {tot_lidos}: O email '{email}' é inválido.")
                continue

            discente, created = Discente.objects.get_or_create(email=email, defaults={'nome': nome})
            if created:
                novos_discentes += 1

            Contato.objects.get_or_create(discente=discente,
                                          defaults={'dtcontato':curso.dt_inicio,
                                                    'origem':'Presencial', 'curso_origem':curso} )

            # Create or update a record in Inscricao
            inscricao, _ = Inscricao.objects.update_or_create(discente=discente, curso=curso, defaults={
                'presenca': tipo != 'matriculas',
                'status': Inscricao.Status.ATIVA
            })

    except Exception as e:
        erros.append(f'Linha {tot_lidos}: Erro durante o processamento: {str(e)}')

    return tot_lidos, novos_discentes, erros


def import_palestrantes(lista, curso):
    tot_lidos = 0
    novos_contatos = 0
    erros = []
    for palestrante in CursoDocente.objects.filter(curso=curso):
        tot_lidos += 1
        nome = palestrante.docente.nome
        email = palestrante.docente.email
        if not email:
            erros.append(f'Docente sem email:{nome}')
            continue
        contato, _ = Contato.objects.get_or_create(email=email, defaults={'nome': nome, 'origem':'Palestrante'})
        novos_contatos += 1
        ListaContato.objects.get_or_create(contato=contato,lista=lista)
    return tot_lidos, novos_contatos, erros


def import_lista(lista, arquivo_xlsx):
    tot_lidos = 0
    tot_invalidos = 0
    novos_contatos = 0
    erros = []
    try:
        wb = load_workbook(arquivo_xlsx)
        ws = wb.active

        for row in ws.iter_rows(min_row=1, values_only=True):
            if len(row) < 1:
                if tot_lidos == 0:
                    erros.append('Arquivo vazio ou com linhas em branco')
                break

            tot_lidos += 1
            nome = row[0]
            if nome:
                nome = nome.strip()
            else:
                erros.append(f"Linha {tot_lidos}: O nome não está preenchido")
                tot_invalidos += 1
                break

            email = row[1]
            if email:
                email = email.strip()
                try:
                    is_valid = validate_email(email, smtp_timeout=10)
                except ValidationError:
                    is_valid = False
            else:
                is_valid = False

            if not is_valid:
                erros.append(f"Linha {tot_lidos}: O email '{email}' é inválido.")
                continue

            contato, created = Contato.objects.get_or_create(email=email, defaults={'nome': nome, 'origem':'Importação'})
            if created:
                novos_contatos += 1

            # Create or update a record in Lista

            ListaContato.objects.update_or_create(contato=contato, lista=lista)

    except Exception as e:
        erros.append(f'Linha {tot_lidos}: Erro durante o processamento: {str(e)}')

    return tot_lidos, novos_contatos, erros
